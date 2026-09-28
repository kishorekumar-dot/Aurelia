import os
import docx

def create_sample_academic_paper(output_path: str):
    doc = docx.Document()
    
    # Set title
    doc.core_properties.title = "A Novel Deep Learning Approach to Quantum Cryptography"
    doc.core_properties.author = "Dr. Alice Smith & Dr. Bob Jones"
    doc.core_properties.subject = "Quantum Computing & Cryptography"
    doc.core_properties.keywords = "quantum cryptography, deep learning, security"
    
    # Document Title (heading 0 isn't always styled standardly, let's use heading 1 or standard title block)
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("A Novel Deep Learning Approach to Quantum Cryptography")
    title_run.bold = True
    title_run.font.size = docx.shared.Pt(24)
    
    author_p = doc.add_paragraph()
    author_run = author_p.add_run("Dr. Alice Smith & Dr. Bob Jones\nDepartment of Computer Science, University of Technology\nAugust 2026")
    author_run.font.size = docx.shared.Pt(11)
    
    # Abstract
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "Quantum cryptography offers theoretically unbreakable encryption based on the laws of physics. "
        "However, practical implementations often suffer from channel noise and detector efficiency issues. "
        "In this paper, we present QCryptNet, a novel deep learning-based framework designed to mitigate physical layer noise "
        "and predict adversarial eavesdropping (Quantum Key Distribution attacks) in real-time. "
        "Our model achieves a 98.4% accuracy in detecting eavesdropping attempts under simulated noisy fibers, "
        "representing a significant step forward in robust, practical quantum secure communication."
    )
    
    # Introduction
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Secure communication is vital in the modern era. Standard public-key cryptographic algorithms (e.g., RSA, ECC) "
        "are vulnerable to quantum computer attacks via Shor's algorithm. Quantum Key Distribution (QKD) is a promising alternative "
        "that uses quantum mechanics to guarantee secure key exchanges. Unfortunately, physical limitations like photon absorption, "
        "polarization drift, and detector dark counts introduce noise that mimics eavesdropping attempts. "
        "Distinguishing noise from active attacks remains a critical challenge."
    )
    doc.add_paragraph(
        "Prior work in this domain has relied heavily on classical statistics or simple thresholding techniques. "
        "These methods fail to capture complex temporal correlation patterns of channel noise. "
        "In this work, we propose QCryptNet, a Recurrent Neural Network (RNN) combined with a Transformer encoder, "
        "to continuously monitor the quantum channel statistics and classify anomalies."
    )
    
    # Related Work
    doc.add_heading("2. Related Work", level=1)
    doc.add_paragraph(
        "Many protocols have been proposed since the original BB84 protocol. "
        "The security of BB84 relies on the fact that any measurement by an eavesdropper (Eve) perturbs the quantum state, "
        "increasing the Quantum Bit Error Rate (QBER). Classical error correction codes are typically employed to correct these errors, "
        "but they do not distinguish the root cause of the error. "
        "Recent machine learning efforts, such as those by Zhao et al. (2023), used Support Vector Machines to classify attacks, "
        "but they lacked real-time adaptation capability."
    )
    
    # Methodology
    doc.add_heading("3. Methodology", level=1)
    doc.add_paragraph(
        "Our proposed system, QCryptNet, consists of three primary modules: the Signal Preprocessing Module, the Temporal Feature Extraction "
        "Module, and the Adversarial Classifier. The signal preprocessing converts raw photon arrival timings and polarization angles "
        "into continuous probability density estimates."
    )
    
    # Subsections
    doc.add_heading("3.1 Neural Network Architecture", level=2)
    doc.add_paragraph(
        "The temporal feature extraction uses a bidirectional LSTM network followed by a multi-head self-attention layer. "
        "This configuration allows the model to process long-term dependencies in the error-rate statistics while prioritizing "
        "burst errors characteristic of active intercept-resend attacks."
    )
    
    # Table of Parameters
    doc.add_heading("3.2 Network Parameters", level=2)
    doc.add_paragraph("Table 1 summarizes the hyperparameters used to train the QCryptNet model.")
    
    table = doc.add_table(rows=4, cols=3)
    # Set header
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Hyperparameter'
    hdr_cells[1].text = 'Value'
    hdr_cells[2].text = 'Description'
    
    # Set data rows
    row1 = table.rows[1].cells
    row1[0].text = 'Learning Rate'
    row1[1].text = '0.0005'
    row1[2].text = 'Adam optimizer starting learning rate'
    
    row2 = table.rows[2].cells
    row2[0].text = 'LSTM Hidden Units'
    row2[1].text = '128'
    row2[2].text = 'Dimensionality of the LSTM hidden state'
    
    row3 = table.rows[3].cells
    row3[0].text = 'Attention Heads'
    row3[1].text = '8'
    row3[2].text = 'Number of attention heads in Transformer encoder'
    
    # Results
    doc.add_heading("4. Evaluation & Results", level=1)
    doc.add_paragraph(
        "We evaluated QCryptNet in a simulated fiber optic environment with variable attenuation coefficients. "
        "We simulated four standard eavesdropping attacks: Intercept-Resend, Photon Number Splitting, "
        "Man-in-the-Middle, and Trojan Horse attacks."
    )
    doc.add_paragraph(
        "Our model outperformed classical threshold-based methods across all scenarios. "
        "Specifically, for low signal-to-noise ratios, QCryptNet successfully detected eavesdropping "
        "with an F1-score of 97.8% compared to only 82.3% using classical methods."
    )
    
    # Conclusion
    doc.add_heading("5. Conclusion & Future Work", level=1)
    doc.add_paragraph(
        "In this study, we introduced QCryptNet, a machine learning solution to secure quantum communications. "
        "By analyzing channel noise signatures in real-time, QCryptNet distinguishes hardware-related losses from malicious interventions. "
        "Future work will focus on deploying this model onto hardware FPGA chips to achieve microsecond inference speeds."
    )
    
    # References
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] C. H. Bennett and G. Brassard, 'Quantum cryptography: Public key distribution and coin tossing,' in Proc. of IEEE International Conference on Computers, Systems and Signal Processing, 1984.")
    doc.add_paragraph("[2] Zhao, Y. et al., 'Machine learning for attack detection in quantum systems,' Journal of Quantum Security, vol. 12, no. 3, 2023.")
    
    # Save the document
    doc.save(output_path)
    print(f"Sample document successfully created at {output_path}")

if __name__ == "__main__":
    os.makedirs(os.path.dirname(output_path := "sample_documents/sample_paper.docx"), exist_ok=True)
    create_sample_academic_paper(output_path)
