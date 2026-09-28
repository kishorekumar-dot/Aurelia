import os
import tempfile
import pytest
import docx
from app.services.document.docx_parser import parse_docx

def test_parse_docx_success():
    # Create a temp docx file
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_name = tmp.name
        
    try:
        # Build document
        doc = docx.Document()
        doc.core_properties.title = "Test Paper Title"
        doc.core_properties.author = "Test Author Name"
        
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("This is paragraph one content.")
        
        # Add table
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Col A"
        table.rows[0].cells[1].text = "Col B"
        table.rows[1].cells[0].text = "Val A"
        table.rows[1].cells[1].text = "Val B"
        
        doc.add_heading("Background", level=2)
        doc.add_paragraph("Subsection paragraph content.")
        
        doc.save(tmp_name)
        
        # Parse it
        result = parse_docx(tmp_name)
        
        # Verify metadata
        assert result["metadata"]["title"] == "Test Paper Title"
        assert result["metadata"]["author"] == "Test Author Name"
        
        # Verify sections
        sections = result["sections"]
        assert len(sections) == 5 # 2 headings, 2 paragraphs, 1 table
        
        assert sections[0]["type"] == "heading"
        assert sections[0]["level"] == 1
        assert sections[0]["text"] == "Introduction"
        assert sections[0]["element_id"] == "HDG-001"
        assert sections[0]["section"] == "Introduction"
        
        assert sections[1]["type"] == "paragraph"
        assert sections[1]["text"] == "This is paragraph one content."
        assert sections[1]["element_id"] == "PAR-001"
        assert sections[1]["section"] == "Introduction"
        
        assert sections[2]["type"] == "table"
        assert sections[2]["rows"] == [["Col A", "Col B"], ["Val A", "Val B"]]
        assert sections[2]["element_id"] == "TBL-001"
        assert sections[2]["section"] == "Introduction"
        
        assert sections[3]["type"] == "heading"
        assert sections[3]["level"] == 2
        assert sections[3]["text"] == "Background"
        assert sections[3]["element_id"] == "HDG-002"
        assert sections[3]["section"] == "Background"
        
        assert sections[4]["type"] == "paragraph"
        assert sections[4]["text"] == "Subsection paragraph content."
        assert sections[4]["element_id"] == "PAR-002"
        assert sections[4]["section"] == "Background"
        
        # Verify markdown output
        assert "# Introduction" in result["markdown"]
        assert "## Background" in result["markdown"]
        assert "| Col A | Col B |" in result["markdown"]
        assert "| --- | --- |" in result["markdown"]
        assert "| Val A | Val B |" in result["markdown"]
        
        # Verify plain text output
        assert "Introduction" in result["full_text"]
        assert "This is paragraph one content." in result["full_text"]
        assert "Col A Col B Val A Val B" in result["full_text"]
        assert "Background" in result["full_text"]
        
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)

def test_parse_docx_invalid_file():
    with pytest.raises(ValueError, match="Failed to open/parse DOCX file"):
        parse_docx("non_existent_file.docx")
