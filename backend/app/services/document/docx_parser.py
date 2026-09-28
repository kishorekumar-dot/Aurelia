import os
from docx import Document
from typing import Dict, Any, List

class DocxParser:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def parse(self) -> Dict[str, Any]:
        return parse_docx(self.file_path)

def parse_docx(file_path: str) -> dict:
    if not os.path.exists(file_path):
        raise ValueError(f"Failed to open/parse DOCX file: {file_path}")
        
    try:
        doc = Document(file_path)
    except Exception as e:
        raise ValueError(f"Failed to open/parse DOCX file: {e}")
        
    metadata = {
        "title": doc.core_properties.title or os.path.basename(file_path),
        "author": doc.core_properties.author or "",
    }
    
    sections = []
    full_text_parts = []
    markdown_parts = []
    
    hdg_count = 0
    par_count = 0
    tbl_count = 0
    
    current_section = "General"
    
    # We iterate over paragraphs and tables in order of document elements
    for block in doc.element.body:
        tag = block.tag.split("}")[-1]
        if tag == "p":
            # Find matching paragraph object
            p = next((p for p in doc.paragraphs if p._element is block), None)
            if not p:
                continue
            text = p.text.strip()
            if not text:
                continue
                
            full_text_parts.append(text)
            
            # Check if heading
            level = 1
            is_heading = False
            if p.style.name.startswith("Heading"):
                try:
                    level = int(p.style.name.replace("Heading", "").strip())
                except Exception:
                    level = 1
                is_heading = True
            
            if is_heading:
                hdg_count += 1
                current_section = text
                element_id = f"HDG-{hdg_count:03d}"
                sections.append({
                    "element_id": element_id,
                    "type": "heading",
                    "text": text,
                    "heading": text,
                    "level": level,
                    "section": current_section,
                    "metadata": {"level": level}
                })
                markdown_parts.append(f"{'#' * level} {text}")
            else:
                par_count += 1
                element_id = f"PAR-{par_count:03d}"
                sections.append({
                    "element_id": element_id,
                    "type": "paragraph",
                    "text": text,
                    "section": current_section,
                    "metadata": {}
                })
                markdown_parts.append(text)
                
        elif tag == "tbl":
            table = next((t for t in doc.tables if t._element is block), None)
            if not table:
                continue
            tbl_count += 1
            element_id = f"TBL-{tbl_count:03d}"
            
            table_rows = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_rows.append(row_data)
                
            sections.append({
                "element_id": element_id,
                "type": "table",
                "rows": table_rows,
                "section": current_section,
                "metadata": {"rows": table_rows}
            })
            
            # Format markdown table
            if table_rows:
                header = table_rows[0]
                markdown_parts.append("| " + " | ".join(header) + " |")
                markdown_parts.append("| " + " | ".join(["---"] * len(header)) + " |")
                for r in table_rows[1:]:
                    markdown_parts.append("| " + " | ".join(r) + " |")
                    
            if table_rows:
                full_text_parts.append(" ".join([" ".join(row) for row in table_rows]))

    full_text_content = "\n\n".join(full_text_parts).strip()
    markdown_content = "\n\n".join(markdown_parts).strip()
    
    return {
        "title": metadata["title"],
        "metadata": metadata,
        "sections": sections,
        "markdown": markdown_content,
        "full_text": full_text_content,
        "estimated_pages": max(1, len(full_text_parts) // 6),
        "references": [line for line in full_text_parts if "reference" in line.lower() or "[" in line]
    }